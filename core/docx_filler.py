"""Fills bracketed placeholders (e.g. ``[Contractor Name]``) in the Contract
for Works template while preserving the original Word formatting.

Only placeholders whose bracket text matches a known field name are
replaced. Any other bracketed text (e.g. stray field-code artifacts such as
``[?]``) is left untouched.
"""

from __future__ import annotations

from docx import Document

from .docx_utils import fill_placeholders, format_date, format_money

# Fields the caller must supply.
REQUIRED_FIELDS = [
    "Project Name",
    "SR Number",
    "Total Amount CHF",
    "Total Amount PH",
    "Contract Sign Date",
    "Contractor Name",
    "Contractor Phone Number",
    "Contractor Email",
    "Contractor Rep Name",
    "Contractor Rep Position",
    "Project Start Date",
    "Project End Date",
    "Total Amount PH in Words",
    "Contractor Account Holder",
    "Contractor Bank Name",
    "Contractor Account Number",
]

DATE_FIELDS = {"Contract Sign Date", "Project Start Date", "Project End Date"}
MONEY_FIELDS = {"Total Amount CHF", "Total Amount PH"}


def prepare_replacements(data: dict) -> dict:
    """Normalizes raw request data into the final placeholder -> text
    map used for substitution, including the derived ``Contractor Rep``
    field (Name, Position)."""
    missing = [f for f in REQUIRED_FIELDS if not str(data.get(f, "")).strip()]
    if missing:
        raise ValueError(f"Missing required field(s): {', '.join(missing)}")

    values = {}
    for field in REQUIRED_FIELDS:
        raw = str(data[field]).strip()
        if field in DATE_FIELDS:
            raw = format_date(raw)
        elif field in MONEY_FIELDS:
            raw = format_money(raw)
        values[field] = raw

    rep_name = values["Contractor Rep Name"]
    rep_position = values["Contractor Rep Position"]
    values["Contractor Rep"] = f"{rep_name}, {rep_position}"

    return values


def fill_template(template_path: str, output_path: str, data: dict) -> str:
    """Fills the Contract for Works template with the given data and writes
    the result to output_path. Returns output_path."""
    replacements = prepare_replacements(data)

    document = Document(template_path)
    fill_placeholders(document, replacements)

    document.save(output_path)
    return output_path
